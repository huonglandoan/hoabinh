import os
import sys
import json
import csv
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from core_excel_utils import (
    apply_header_style,
    apply_data_style,
    apply_total_style,
    autofit_column_widths,
    FMT_CURRENCY,
    FMT_DATE,
    FMT_PERCENT,
    COLOR_ZEBRA_BG,
    COLOR_TOTAL_BG
)
from core_excel_utils import vn_slug

# PDF Font registration
try:
    font_dir = "/System/Library/Fonts/Supplemental"
    pdfmetrics.registerFont(TTFont('Arial', os.path.join(font_dir, 'Arial.ttf')))
    pdfmetrics.registerFont(TTFont('Arial-Bold', os.path.join(font_dir, 'Arial Bold.ttf')))
    pdfmetrics.registerFont(TTFont('Arial-Italic', os.path.join(font_dir, 'Arial Italic.ttf')))
    FONT = 'Arial'
except Exception:
    FONT = 'Helvetica'

def format_currency(val):
    try:
        return f"{int(val):,}".replace(",", ".")
    except (ValueError, TypeError):
        return str(val)

def get_display_name(item):
    name = item.get("item_name", "")
    if item.get("is_extra") or item.get("isExtra"):
        if not name.endswith("(Phát sinh)"):
            name += " (Phát sinh)"
    return name


def generate_payment_package(input_data_path, project_id=None):
    with open(input_data_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    billing_info = data.get("billing_info", {})
    items = data.get("items", [])
    invoices = data.get("invoices", [])
    
    project_name = billing_info.get("project_name", "Cong_Trinh")
    project_slug = vn_slug(project_name)
    
    project_id = project_id or project_slug
    
    # Paths (absolute paths)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(script_dir))
    data_dir = os.path.join(project_root, "data")
    
    master_dir = os.path.join(data_dir, "master", project_id)
    exports_dir = os.path.join(data_dir, "exports", project_id, "payments")
    os.makedirs(master_dir, exist_ok=True)
    os.makedirs(exports_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    excel_filename = f"Ho_So_Thanh_Toan_{project_slug}_{timestamp}.xlsx"
    excel_path = os.path.join(exports_dir, excel_filename)
    master_excel_path = os.path.join(master_dir, "payment_profile.xlsx")
    
    pdf_filename = f"Ho_So_Thanh_Toan_{project_slug}_{timestamp}.pdf"
    pdf_path = os.path.join(exports_dir, pdf_filename)
    master_pdf_path = os.path.join(master_dir, "payment_profile.pdf")
    
    csv_filename = f"Bang_Ke_HDGTGT_{project_slug}_{timestamp}.csv"
    csv_path = os.path.join(exports_dir, csv_filename)
    master_csv_path = os.path.join(master_dir, "invoices_data.csv")
    
    # 1. Build Excel
    build_excel_sheets(data, excel_path)
    
    # 2. Build PDF
    build_combined_pdf(data, pdf_path)
    
    # 3. Build CSV
    build_invoices_csv(invoices, csv_path)
    
    # Copy to master
    import shutil
    # Backup existing master files
    try:
        backups_dir = os.path.join(project_root, 'data', 'backups', project_id, 'payments')
        os.makedirs(backups_dir, exist_ok=True)
        if os.path.exists(master_excel_path):
            shutil.copy2(master_excel_path, os.path.join(backups_dir, os.path.basename(master_excel_path) + f".backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}"))
        if os.path.exists(master_pdf_path):
            shutil.copy2(master_pdf_path, os.path.join(backups_dir, os.path.basename(master_pdf_path) + f".backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}"))
        if os.path.exists(master_csv_path):
            shutil.copy2(master_csv_path, os.path.join(backups_dir, os.path.basename(master_csv_path) + f".backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}"))
    except Exception:
        pass
    shutil.copy2(excel_path, master_excel_path)
    shutil.copy2(pdf_path, master_pdf_path)
    shutil.copy2(csv_path, master_csv_path)
    
    return {
        "output_excel_path": excel_path,
        "output_pdf_path": pdf_path,
        "output_csv_path": csv_path,
        "master_excel_path": master_excel_path,
        "master_pdf_path": master_pdf_path,
        "master_csv_path": master_csv_path,
        "project_id": project_id
    }

def build_excel_sheets(data, out_path):
    wb = openpyxl.Workbook()
    
    billing_info = data.get("billing_info", {})
    items = data.get("items", [])
    invoices = data.get("invoices", [])
    
    # Common styles
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)
    
    # --- 1. Sheet BAO_GIA ---
    ws1 = wb.active
    ws1.title = "BAO_GIA"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1['A1'] = billing_info.get("contractor_name", "Bên B")
    ws1['F1'] = f"Số HĐ: {billing_info.get('contract_code', '')}"
    ws1['F2'] = f"Ngày HĐ: {billing_info.get('contract_date', '')}"
    
    ws1['A4'] = "BẢNG BÁO GIÁ CHI TIẾT HỢP ĐỒNG"
    ws1['A4'].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    ws1['A4'].alignment = Alignment(horizontal="center")
    ws1.merge_cells("A4:F4")
    
    ws1['A6'] = "Dự án:"
    ws1['B6'] = billing_info.get("project_name", "")
    ws1['A7'] = "Bên A:"
    ws1['B7'] = billing_info.get("client_name", "")
    ws1['D7'] = "Bên B:"
    ws1['E7'] = billing_info.get("contractor_name", "")
    
    for r in range(6, 8):
        ws1[f'A{r}'].font = font_bold
        ws1[f'B{r}'].font = font_regular
        if ws1[f'D{r}'].coordinate:
            ws1[f'D{r}'].font = font_bold
            ws1[f'E{r}'].font = font_regular
            
    # Headers
    headers_bg = ["STT", "Hạng mục công việc", "ĐVT", "KL Dự toán", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"]
    for col_idx, text in enumerate(headers_bg, 1):
        cell = ws1.cell(row=9, column=col_idx, value=text)
        apply_header_style(cell)
        
    r_idx = 10
    for idx, item in enumerate(items, 1):
        ws1.cell(row=r_idx, column=1, value=idx)
        ws1.cell(row=r_idx, column=2, value=get_display_name(item))
        ws1.cell(row=r_idx, column=3, value=item.get("unit", ""))
        ws1.cell(row=r_idx, column=4, value=float(item.get("contract_quantity") or 0))
        ws1.cell(row=r_idx, column=5, value=float(item.get("unit_price") or 0))
        ws1.cell(row=r_idx, column=6, value=f"=D{r_idx}*E{r_idx}")
        
        # Styles
        bg_zebra = COLOR_ZEBRA_BG if idx % 2 == 0 else None
        for c in range(1, 7):
            cell = ws1.cell(row=r_idx, column=c)
            apply_data_style(cell, bg_color=bg_zebra)
            if c in (1, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c in (4, 5, 6):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = FMT_CURRENCY
        r_idx += 1
        
    # Totals in BAO_GIA
    # Tổng cộng
    ws1.merge_cells(f"A{r_idx}:E{r_idx}")
    ws1.cell(row=r_idx, column=1, value="TỔNG CỘNG").alignment = Alignment(horizontal="right")
    apply_total_style(ws1.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws1.cell(row=r_idx, column=c))
    ws1.cell(row=r_idx, column=6, value=f"=SUM(F10:F{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws1.cell(row=r_idx, column=6), align="right")
    
    # QLCP (5%)
    r_idx += 1
    ws1.merge_cells(f"A{r_idx}:E{r_idx}")
    ws1.cell(row=r_idx, column=1, value="Chi phí quản lý (5%)").alignment = Alignment(horizontal="right")
    apply_total_style(ws1.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws1.cell(row=r_idx, column=c))
    ws1.cell(row=r_idx, column=6, value=f"=F{r_idx-1}*0.05").number_format = FMT_CURRENCY
    apply_total_style(ws1.cell(row=r_idx, column=6), align="right")
    
    # VAT (10%)
    r_idx += 1
    ws1.merge_cells(f"A{r_idx}:E{r_idx}")
    ws1.cell(row=r_idx, column=1, value="Thuế VAT (10%)").alignment = Alignment(horizontal="right")
    apply_total_style(ws1.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws1.cell(row=r_idx, column=c))
    ws1.cell(row=r_idx, column=6, value=f"=(F{r_idx-2}+F{r_idx-1})*0.10").number_format = FMT_CURRENCY
    apply_total_style(ws1.cell(row=r_idx, column=6), align="right")
    
    # Tổng giá trị hợp đồng
    r_idx += 1
    ws1.merge_cells(f"A{r_idx}:E{r_idx}")
    ws1.cell(row=r_idx, column=1, value="TỔNG GIÁ TRỊ HỢP ĐỒNG").alignment = Alignment(horizontal="right")
    apply_total_style(ws1.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws1.cell(row=r_idx, column=c))
    ws1.cell(row=r_idx, column=6, value=f"=F{r_idx-3}+F{r_idx-2}+F{r_idx-1}").number_format = FMT_CURRENCY
    apply_total_style(ws1.cell(row=r_idx, column=6), align="right")
    
    autofit_column_widths(ws1, padding=3, manual_widths={'A': 5, 'B': 40, 'C': 8, 'D': 12, 'E': 15, 'F': 18})

    # --- 2. Sheet NGHIEM_THU ---
    ws2 = wb.create_sheet(title="NGHIEM_THU")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2['A1'] = billing_info.get("contractor_name", "Bên B")
    ws2['H1'] = f"Căn cứ HĐ số: {billing_info.get('contract_code', '')}"
    
    ws2['A3'] = "BIÊN BẢN NGHIỆM THU KHỐI LƯỢNG HOÀN THÀNH"
    ws2['A3'].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    ws2['A3'].alignment = Alignment(horizontal="center")
    ws2.merge_cells("A3:H3")
    
    ws2['A5'] = "Dự án:"
    ws2['B5'] = billing_info.get("project_name", "")
    
    headers_nt = ["STT", "Hạng mục công việc", "ĐVT", "KL theo HĐ", "KL thực hiện", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)", "Ghi chú"]
    for col_idx, text in enumerate(headers_nt, 1):
        cell = ws2.cell(row=7, column=col_idx, value=text)
        apply_header_style(cell)
        
    r_idx = 8
    for idx, item in enumerate(items, 1):
        ws2.cell(row=r_idx, column=1, value=idx)
        ws2.cell(row=r_idx, column=2, value=get_display_name(item))
        ws2.cell(row=r_idx, column=3, value=item.get("unit", ""))
        
        # Link contract quantity to BAO_GIA
        ws2.cell(row=r_idx, column=4, value=f"=BAO_GIA!D{9+idx}")
        # Actual quantity
        ws2.cell(row=r_idx, column=5, value=float(item.get("actual_quantity") or 0))
        # Unit price links to BAO_GIA
        ws2.cell(row=r_idx, column=6, value=f"=BAO_GIA!E{9+idx}")
        # Thành tiền = KL thực hiện * Đơn giá
        ws2.cell(row=r_idx, column=7, value=f"=E{r_idx}*F{r_idx}")
        
        # Check notes
        ws2.cell(row=r_idx, column=8, value=f'=IF(E{r_idx}>=D{r_idx}, "Đạt", "Thiếu " & TEXT(D{r_idx}-E{r_idx}, "#,##0"))')
        
        bg_zebra = COLOR_ZEBRA_BG if idx % 2 == 0 else None
        for c in range(1, 9):
            cell = ws2.cell(row=r_idx, column=c)
            apply_data_style(cell, bg_color=bg_zebra)
            if c in (1, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c in (4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = FMT_CURRENCY
        r_idx += 1
        
    # Totals in NGHIEM_THU
    ws2.merge_cells(f"A{r_idx}:F{r_idx}")
    ws2.cell(row=r_idx, column=1, value="TỔNG CỘNG KHỐI LƯỢNG NGHIỆM THU").alignment = Alignment(horizontal="right")
    apply_total_style(ws2.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws2.cell(row=r_idx, column=c))
    ws2.cell(row=r_idx, column=7, value=f"=SUM(G8:G{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws2.cell(row=r_idx, column=7), align="right")
    apply_total_style(ws2.cell(row=r_idx, column=8))
    
    autofit_column_widths(ws2, padding=3, manual_widths={'A': 5, 'B': 40, 'C': 8, 'D': 12, 'E': 12, 'F': 15, 'G': 18, 'H': 12})

    # --- 3. Sheet QUYET_TOAN_AB ---
    ws3 = wb.create_sheet(title="QUYET_TOAN_AB")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3['A1'] = "BIÊN BẢN QUYẾT TOÁN HỢP ĐỒNG (A-B)"
    ws3['A1'].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    ws3['A1'].alignment = Alignment(horizontal="center")
    ws3.merge_cells("A1:H1")
    
    headers_qt = ["STT", "Hạng mục", "KL HĐ", "Đơn giá HĐ", "GT Hợp đồng", "KL Thực hiện", "GT Thực hiện", "Chênh lệch"]
    for col_idx, text in enumerate(headers_qt, 1):
        cell = ws3.cell(row=3, column=col_idx, value=text)
        apply_header_style(cell)
        
    r_idx = 4
    for idx, item in enumerate(items, 1):
        ws3.cell(row=r_idx, column=1, value=idx)
        ws3.cell(row=r_idx, column=2, value=get_display_name(item))
        
        # Link contract quantity, price, contract amount from BAO_GIA
        ws3.cell(row=r_idx, column=3, value=f"=BAO_GIA!D{9+idx}")
        ws3.cell(row=r_idx, column=4, value=f"=BAO_GIA!E{9+idx}")
        ws3.cell(row=r_idx, column=5, value=f"=C{r_idx}*D{r_idx}")
        
        # Link actual quantity from NGHIEM_THU, and price from BAO_GIA
        ws3.cell(row=r_idx, column=6, value=f"=NGHIEM_THU!E{7+idx}")
        ws3.cell(row=r_idx, column=7, value=f"=F{r_idx}*D{r_idx}")
        
        # Chênh lệch = GT Thực hiện - GT Hợp đồng
        ws3.cell(row=r_idx, column=8, value=f"=G{r_idx}-E{r_idx}")
        
        bg_zebra = COLOR_ZEBRA_BG if idx % 2 == 0 else None
        for c in range(1, 9):
            cell = ws3.cell(row=r_idx, column=c)
            apply_data_style(cell, bg_color=bg_zebra)
            if c == 1:
                cell.alignment = Alignment(horizontal="center")
            elif c in (3, 4, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = FMT_CURRENCY
        r_idx += 1
        
    # Totals in QUYET_TOAN_AB
    ws3.merge_cells(f"A{r_idx}:D{r_idx}")
    ws3.cell(row=r_idx, column=1, value="TỔNG CỘNG").alignment = Alignment(horizontal="right")
    apply_total_style(ws3.cell(row=r_idx, column=1))
    for c in range(2, 5): apply_total_style(ws3.cell(row=r_idx, column=c))
    
    # Contract total
    ws3.cell(row=r_idx, column=5, value=f"=SUM(E4:E{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws3.cell(row=r_idx, column=5), align="right")
    
    # Actual total quantity cell is blank
    apply_total_style(ws3.cell(row=r_idx, column=6))
    
    # Actual total value
    ws3.cell(row=r_idx, column=7, value=f"=SUM(G4:G{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws3.cell(row=r_idx, column=7), align="right")
    
    # Total variance
    ws3.cell(row=r_idx, column=8, value=f"=G{r_idx}-E{r_idx}").number_format = FMT_CURRENCY
    apply_total_style(ws3.cell(row=r_idx, column=8), align="right")
    
    autofit_column_widths(ws3, padding=3, manual_widths={'A': 5, 'B': 40, 'C': 12, 'D': 15, 'E': 18, 'F': 12, 'G': 18, 'H': 18})

    # --- 4. Sheet DE_NGHI_TT ---
    ws4 = wb.create_sheet(title="DE_NGHI_TT")
    ws4.views.sheetView[0].showGridLines = True
    
    ws4['A1'] = "ĐỀ NGHỊ THANH TOÁN"
    ws4['A1'].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    
    ws4['A3'] = "Kính gửi:"
    ws4['B3'] = billing_info.get("client_name", "")
    ws4['A4'] = "Dự án:"
    ws4['B4'] = billing_info.get("project_name", "")
    ws4['A5'] = "Bên đề nghị:"
    ws4['B5'] = billing_info.get("contractor_name", "")
    ws4['A6'] = "Hợp đồng số:"
    ws4['B6'] = f"{billing_info.get('contract_code', '')} ký ngày {billing_info.get('contract_date', '')}"
    
    for r in range(3, 7):
        ws4[f'A{r}'].font = font_bold
        ws4[f'B{r}'].font = font_regular
        
    # Table headers starting at row 9
    headers_tt = ["Nội dung thanh toán", "Giá trị (VNĐ)"]
    ws4.merge_cells("A9:C9")
    ws4.cell(row=9, column=1, value="Nội dung thanh toán")
    ws4.cell(row=9, column=4, value="Giá trị (VNĐ)")
    apply_header_style(ws4.cell(row=9, column=1))
    apply_header_style(ws4.cell(row=9, column=4))
    
    # Dòng 1: Giá trị khối lượng hoàn thành theo hợp đồng (lấy từ quyết toán)
    ws4.cell(row=10, column=1, value="1. Giá trị KL hoàn thành thực tế (trước thuế)")
    ws4.merge_cells("A10:C10")
    # Link to QUYET_TOAN_AB total actual
    ws4.cell(row=10, column=4, value=f"=QUYET_TOAN_AB!G{r_idx}").number_format = FMT_CURRENCY
    
    # Dòng 2: Chi phí phát sinh (nếu có)
    ws4.cell(row=11, column=1, value="2. Giá trị KL công việc phát sinh")
    ws4.merge_cells("A11:C11")
    ws4.cell(row=11, column=4, value=float(billing_info.get("additional_value") or 0)).number_format = FMT_CURRENCY
    
    # Dòng 3: Tạm ứng khấu trừ
    ws4.cell(row=12, column=1, value="3. Giảm trừ tiền tạm ứng đã nhận")
    ws4.merge_cells("A12:C12")
    # Store as a negative number
    adv = float(billing_info.get("advance_deduction") or 0)
    ws4.cell(row=12, column=4, value=-abs(adv)).number_format = FMT_CURRENCY
    
    # Dòng 4: Cộng trước VAT
    ws4.cell(row=13, column=1, value="4. Cộng giá trị đề nghị thanh toán (trước VAT)")
    ws4.merge_cells("A13:C13")
    ws4.cell(row=13, column=4, value="=SUM(D10:D12)").number_format = FMT_CURRENCY
    
    # Dòng 5: Thuế VAT
    ws4.cell(row=14, column=1, value="5. Thuế GTGT (10%)")
    ws4.merge_cells("A14:C14")
    ws4.cell(row=14, column=4, value="=D13*0.1").number_format = FMT_CURRENCY
    
    # Dòng 6: Tổng cộng đề nghị thanh toán
    ws4.cell(row=15, column=1, value="6. TỔNG CỘNG TIỀN ĐỀ NGHỊ THANH TOÁN (SAU VAT)")
    ws4.merge_cells("A15:C15")
    ws4.cell(row=15, column=4, value="=D13+D14").number_format = FMT_CURRENCY
    
    # Apply data styles
    for row in range(10, 16):
        bg = COLOR_ZEBRA_BG if row % 2 == 0 else None
        # Format label cells
        for c in range(1, 4):
            apply_data_style(ws4.cell(row=row, column=c), bg_color=bg)
        # Format amount cell
        cell_val = ws4.cell(row=row, column=4)
        if row == 15:
            apply_total_style(cell_val, align="right", num_format=FMT_CURRENCY)
            # Make label bold too
            for c in range(1, 4):
                ws4.cell(row=row, column=c).font = Font(name="Arial", size=10, bold=True)
                ws4.cell(row=row, column=c).fill = PatternFill(start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type="solid")
        else:
            apply_data_style(cell_val, align="right", num_format=FMT_CURRENCY, bg_color=bg)
            
    # Banking details
    ws4['A17'] = "Hình thức thanh toán: Chuyển khoản"
    ws4['A17'].font = font_bold
    ws4['A18'] = f"Tài khoản thụ hưởng: {billing_info.get('bank_account', '')} tại {billing_info.get('bank_name', '')}"
    ws4['A18'].font = font_regular
    
    ws4.column_dimensions['A'].width = 15
    ws4.column_dimensions['B'].width = 15
    ws4.column_dimensions['C'].width = 25
    ws4.column_dimensions['D'].width = 20

    # --- 5. Sheet BANG_KE_HDGTGT ---
    ws5 = wb.create_sheet(title="BANG_KE_HDGTGT")
    ws5.views.sheetView[0].showGridLines = True
    
    ws5['A1'] = "BẢNG KÊ HÓA ĐƠN GIÁ TRỊ GIA TĂNG (GTGT)"
    ws5['A1'].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    ws5['A1'].alignment = Alignment(horizontal="center")
    ws5.merge_cells("A1:I1")
    
    headers_inv = ["STT", "Số hóa đơn", "Ngày hóa đơn", "Tên người bán/Nhà cung cấp", "Loại chi phí", "Trước thuế (VNĐ)", "Thuế suất", "Tiền thuế (VNĐ)", "Tổng thanh toán (VNĐ)"]
    for col_idx, text in enumerate(headers_inv, 1):
        cell = ws5.cell(row=4, column=col_idx, value=text)
        apply_header_style(cell)
        
    r_idx = 5
    for idx, inv in enumerate(invoices, 1):
        ws5.cell(row=r_idx, column=1, value=idx)
        ws5.cell(row=r_idx, column=2, value=inv.get("so_hd", ""))
        ws5.cell(row=r_idx, column=3, value=inv.get("ngay", ""))
        ws5.cell(row=r_idx, column=4, value=inv.get("ten_ncc", ""))
        ws5.cell(row=r_idx, column=5, value=inv.get("loai", "Vật tư"))
        ws5.cell(row=r_idx, column=6, value=float(inv.get("truoc_thue") or 0))
        ws5.cell(row=r_idx, column=7, value=float(inv.get("thue_suat") or 0.10))
        # Formulas
        ws5.cell(row=r_idx, column=8, value=f"=F{r_idx}*G{r_idx}")
        ws5.cell(row=r_idx, column=9, value=f"=F{r_idx}+H{r_idx}")
        
        bg_zebra = COLOR_ZEBRA_BG if idx % 2 == 0 else None
        for c in range(1, 10):
            cell = ws5.cell(row=r_idx, column=c)
            apply_data_style(cell, bg_color=bg_zebra)
            if c in (1, 2, 3, 5):
                cell.alignment = Alignment(horizontal="center")
            elif c in (6, 8, 9):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = FMT_CURRENCY
            elif c == 7:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = FMT_PERCENT
        r_idx += 1
        
    # Totals in BANG_KE_HDGTGT
    ws5.merge_cells(f"A{r_idx}:E{r_idx}")
    ws5.cell(row=r_idx, column=1, value="TỔNG CỘNG").alignment = Alignment(horizontal="right")
    apply_total_style(ws5.cell(row=r_idx, column=1))
    for c in range(2, 6): apply_total_style(ws5.cell(row=r_idx, column=c))
    
    # Sum trước thuế
    ws5.cell(row=r_idx, column=6, value=f"=SUM(F5:F{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws5.cell(row=r_idx, column=6), align="right")
    # Empty for tax rate
    apply_total_style(ws5.cell(row=r_idx, column=7))
    # Sum tiền thuế
    ws5.cell(row=r_idx, column=8, value=f"=SUM(H5:H{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws5.cell(row=r_idx, column=8), align="right")
    # Sum tổng thanh toán
    ws5.cell(row=r_idx, column=9, value=f"=SUM(I5:I{r_idx-1})").number_format = FMT_CURRENCY
    apply_total_style(ws5.cell(row=r_idx, column=9), align="right")
    
    autofit_column_widths(ws5, padding=3, manual_widths={'A': 5, 'B': 12, 'C': 12, 'D': 35, 'E': 12, 'F': 18, 'G': 10, 'H': 15, 'I': 18})
    
    wb.save(out_path)

def build_combined_pdf(data, out_path):
    # Combined PDF in Portrait A4
    doc = SimpleDocTemplate(
        out_path,
        pagesize=A4,
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=15*mm,
        bottomMargin=15*mm
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle', parent=styles['Heading1'],
        fontName=f"{FONT}-Bold" if FONT == 'Arial' else 'Helvetica-Bold',
        fontSize=15, textColor=colors.HexColor('#1F4E79'),
        alignment=1, spaceAfter=15
    )
    
    section_style = ParagraphStyle(
        'DocSection', parent=styles['Heading2'],
        fontName=f"{FONT}-Bold" if FONT == 'Arial' else 'Helvetica-Bold',
        fontSize=11, textColor=colors.HexColor('#1F4E79'),
        spaceBefore=10, spaceAfter=5
    )
    
    info_style = ParagraphStyle(
        'InfoText', parent=styles['Normal'],
        fontName=FONT, fontSize=9, leading=13,
        textColor=colors.HexColor('#333333')
    )
    
    cell_style = ParagraphStyle(
        'CellText', parent=styles['Normal'],
        fontName=FONT, fontSize=8, leading=10
    )
    cell_center = ParagraphStyle('CellTextCenter', parent=cell_style, alignment=1)
    cell_right = ParagraphStyle('CellTextRight', parent=cell_style, alignment=2)
    cell_header = ParagraphStyle('CellTextHeader', parent=cell_style, fontName=f"{FONT}-Bold" if FONT == 'Arial' else 'Helvetica-Bold', textColor=colors.white, alignment=1)
    cell_total = ParagraphStyle('CellTextTotal', parent=cell_style, fontName=f"{FONT}-Bold" if FONT == 'Arial' else 'Helvetica-Bold', alignment=2)

    story = []
    
    billing_info = data.get("billing_info", {})
    items = data.get("items", [])
    invoices = data.get("invoices", [])
    
    # ------------------ Section 1: ĐỀ NGHỊ THANH TOÁN ------------------
    story.append(Paragraph("ĐỀ NGHỊ THANH TOÁN", title_style))
    story.append(Spacer(1, 5))
    
    info_text = f"""
    <b>Kính gửi:</b> {billing_info.get('client_name', '')}<br/>
    <b>Dự án:</b> {billing_info.get('project_name', '')}<br/>
    <b>Bên đề nghị:</b> {billing_info.get('contractor_name', '')}<br/>
    <b>Căn cứ Hợp đồng số:</b> {billing_info.get('contract_code', '')} ký ngày {billing_info.get('contract_date', '')}<br/>
    """
    story.append(Paragraph(info_text, info_style))
    story.append(Spacer(1, 10))
    
    # Calc values for display in PDF
    total_val = sum(float(x.get("actual_quantity", 0)) * float(x.get("unit_price", 0)) for x in items)
    add_val = float(billing_info.get("additional_value", 0))
    adv_ded = float(billing_info.get("advance_deduction", 0))
    sub_total = total_val + add_val - adv_ded
    vat_val = sub_total * 0.10
    grand_total = sub_total + vat_val
    
    tt_headers = [Paragraph("Nội dung thanh toán", cell_header), Paragraph("Giá trị (VNĐ)", cell_header)]
    tt_rows = [
        tt_headers,
        [Paragraph("1. Giá trị KL hoàn thành thực tế (trước thuế)", cell_style), Paragraph(format_currency(total_val), cell_right)],
        [Paragraph("2. Giá trị KL công việc phát sinh", cell_style), Paragraph(format_currency(add_val), cell_right)],
        [Paragraph("3. Giảm trừ tiền tạm ứng đã nhận", cell_style), Paragraph(f"- {format_currency(adv_ded)}" if adv_ded > 0 else "0", cell_right)],
        [Paragraph("4. Cộng giá trị đề nghị thanh toán (trước VAT)", cell_style), Paragraph(format_currency(sub_total), cell_right)],
        [Paragraph("5. Thuế GTGT (10%)", cell_style), Paragraph(format_currency(vat_val), cell_right)],
        [Paragraph("<b>6. TỔNG CỘNG TIỀN ĐỀ NGHỊ THANH TOÁN (SAU VAT)</b>", cell_style), Paragraph(f"<b>{format_currency(grand_total)}</b>", cell_right)]
    ]
    
    tt_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D3D3D3')),
        ('BACKGROUND', (0, 4), (-1, 4), colors.HexColor('#EBF3FB')),
        ('BACKGROUND', (0, 6), (-1, 6), colors.HexColor('#BDD7EE')),
    ]
    tt_table = Table(tt_rows, colWidths=[120*mm, 60*mm])
    tt_table.setStyle(TableStyle(tt_styles))
    story.append(tt_table)
    
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"<b>Hình thức thanh toán:</b> Chuyển khoản<br/><b>Tài khoản thụ hưởng:</b> {billing_info.get('bank_account', '')} tại {billing_info.get('bank_name', '')}", info_style))
    
    # Signatures
    story.append(Spacer(1, 15))
    sig_data = [
        [
            Paragraph("<b>Người lập</b><br/><i>(Ký tên)</i>", cell_center),
            Paragraph("<b>Kế toán trưởng</b><br/><i>(Ký tên)</i>", cell_center),
            Paragraph("<b>Đại diện Bên B</b><br/><i>(Ký tên, đóng dấu)</i>", cell_center)
        ]
    ]
    sig_table = Table(sig_data, colWidths=[60*mm, 60*mm, 60*mm])
    sig_table.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER')]))
    story.append(sig_table)
    
    # ------------------ Section 2: BÁO GIÁ CHI TIẾT ------------------
    story.append(PageBreak())
    story.append(Paragraph("BẢNG BÁO GIÁ CHI TIẾT HỢP ĐỒNG", title_style))
    story.append(Spacer(1, 5))
    
    bg_headers = [
        Paragraph("STT", cell_header), Paragraph("Hạng mục công việc", cell_header),
        Paragraph("ĐVT", cell_header), Paragraph("Số lượng", cell_header),
        Paragraph("Đơn giá (VND)", cell_header), Paragraph("Thành tiền (VND)", cell_header)
    ]
    bg_rows = [bg_headers]
    bg_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D3D3D3')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]
    
    tot_contract = 0
    for idx, item in enumerate(items, 1):
        qty = float(item.get("contract_quantity") or 0)
        prc = float(item.get("unit_price") or 0)
        amt = qty * prc
        tot_contract += amt
        
        bg = colors.white if idx % 2 != 0 else colors.HexColor('#EBF3FB')
        bg_rows.append([
            Paragraph(str(idx), cell_center), Paragraph(get_display_name(item), cell_style),
            Paragraph(item.get("unit", ""), cell_center), Paragraph(format_currency(qty), cell_right),
            Paragraph(format_currency(prc), cell_right), Paragraph(format_currency(amt), cell_right)
        ])
        for c in range(6):
            bg_styles.append(('BACKGROUND', (c, idx), (c, idx), bg))
            
    # Add summary in BAO_GIA
    tot_row = len(items) + 1
    bg_rows.append([Paragraph("<b>Cộng giá trị dự toán</b>", cell_total)] + [Paragraph("", cell_style)]*4 + [Paragraph(format_currency(tot_contract), cell_right)])
    bg_rows.append([Paragraph("Chi phí quản lý (5%)", cell_total)] + [Paragraph("", cell_style)]*4 + [Paragraph(format_currency(tot_contract*0.05), cell_right)])
    bg_rows.append([Paragraph("Thuế VAT (10%)", cell_total)] + [Paragraph("", cell_style)]*4 + [Paragraph(format_currency((tot_contract*1.05)*0.1), cell_right)])
    bg_rows.append([Paragraph("<b>TỔNG GIÁ TRỊ HỢP ĐỒNG</b>", cell_total)] + [Paragraph("", cell_style)]*4 + [Paragraph(format_currency(tot_contract*1.05*1.1), cell_right)])
    
    for r in range(tot_row, tot_row + 4):
        bg_styles.append(('SPAN', (0, r), (4, r)))
        bg_styles.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#BDD7EE') if r == tot_row + 3 else colors.HexColor('#EBF3FB')))
        
    bg_table = Table(bg_rows, colWidths=[10*mm, 85*mm, 15*mm, 20*mm, 25*mm, 25*mm])
    bg_table.setStyle(TableStyle(bg_styles))
    story.append(bg_table)
    
    # ------------------ Section 3: BIÊN BẢN NGHIỆM THU ------------------
    story.append(PageBreak())
    story.append(Paragraph("BIÊN BẢN NGHIỆM THU KHỐI LƯỢNG HOÀN THÀNH", title_style))
    story.append(Spacer(1, 5))
    
    nt_headers = [
        Paragraph("STT", cell_header), Paragraph("Hạng mục công việc", cell_header),
        Paragraph("ĐVT", cell_header), Paragraph("KL HĐ", cell_header),
        Paragraph("KL Thực", cell_header), Paragraph("Đơn giá", cell_header),
        Paragraph("Thành tiền (VND)", cell_header)
    ]
    nt_rows = [nt_headers]
    nt_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D3D3D3')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]
    
    tot_actual = 0
    for idx, item in enumerate(items, 1):
        c_qty = float(item.get("contract_quantity") or 0)
        a_qty = float(item.get("actual_quantity") or 0)
        prc = float(item.get("unit_price") or 0)
        amt = a_qty * prc
        tot_actual += amt
        
        bg = colors.white if idx % 2 != 0 else colors.HexColor('#EBF3FB')
        nt_rows.append([
            Paragraph(str(idx), cell_center), Paragraph(get_display_name(item), cell_style),
            Paragraph(item.get("unit", ""), cell_center), Paragraph(format_currency(c_qty), cell_right),
            Paragraph(format_currency(a_qty), cell_right), Paragraph(format_currency(prc), cell_right),
            Paragraph(format_currency(amt), cell_right)
        ])
        for c in range(7):
            nt_styles.append(('BACKGROUND', (c, idx), (c, idx), bg))
            
    tot_row = len(items) + 1
    nt_rows.append([Paragraph("<b>TỔNG CỘNG NGHIỆM THU (TRƯỚC THUẾ)</b>", cell_total)] + [Paragraph("", cell_style)]*5 + [Paragraph(format_currency(tot_actual), cell_right)])
    nt_styles.append(('SPAN', (0, tot_row), (5, tot_row)))
    nt_styles.append(('BACKGROUND', (0, tot_row), (-1, tot_row), colors.HexColor('#BDD7EE')))
    
    nt_table = Table(nt_rows, colWidths=[10*mm, 75*mm, 12*mm, 20*mm, 20*mm, 20*mm, 23*mm])
    nt_table.setStyle(TableStyle(nt_styles))
    story.append(nt_table)
    
    story.append(Spacer(1, 10))
    story.append(Paragraph("<b>Kết luận:</b> Hai bên thống nhất nghiệm thu khối lượng công việc hoàn thành đạt yêu cầu kỹ thuật và đủ điều kiện thanh toán.", info_style))
    
    story.append(Spacer(1, 15))
    sig_data_nt = [
        [
            Paragraph("<b>Người lập</b><br/><i>(Ký tên)</i>", cell_center),
            Paragraph("<b>Kỹ thuật Bên B</b><br/><i>(Ký tên)</i>", cell_center),
            Paragraph("<b>Đại diện Bên A</b><br/><i>(Ký tên, đóng dấu)</i>", cell_center)
        ]
    ]
    sig_table_nt = Table(sig_data_nt, colWidths=[60*mm, 60*mm, 60*mm])
    sig_table_nt.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER')]))
    story.append(sig_table_nt)
    
    # ------------------ Section 4: QUYẾT TOÁN A-B ------------------
    story.append(PageBreak())
    story.append(Paragraph("BIÊN BẢN QUYẾT TOÁN HỢP ĐỒNG (A-B)", title_style))
    story.append(Spacer(1, 5))
    
    qt_headers = [
        Paragraph("STT", cell_header), Paragraph("Hạng mục công việc", cell_header),
        Paragraph("KL HĐ", cell_header), Paragraph("GT Hợp đồng", cell_header),
        Paragraph("KL Thực tế", cell_header), Paragraph("GT Thực tế", cell_header),
        Paragraph("Chênh lệch (VND)", cell_header)
    ]
    qt_rows = [qt_headers]
    qt_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D3D3D3')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]
    
    tot_con_val = 0
    tot_act_val = 0
    for idx, item in enumerate(items, 1):
        c_qty = float(item.get("contract_quantity") or 0)
        a_qty = float(item.get("actual_quantity") or 0)
        prc = float(item.get("unit_price") or 0)
        c_val = c_qty * prc
        a_val = a_qty * prc
        diff = a_val - c_val
        tot_con_val += c_val
        tot_act_val += a_val
        
        bg = colors.white if idx % 2 != 0 else colors.HexColor('#EBF3FB')
        qt_rows.append([
            Paragraph(str(idx), cell_center), Paragraph(get_display_name(item), cell_style),
            Paragraph(format_currency(c_qty), cell_right), Paragraph(format_currency(c_val), cell_right),
            Paragraph(format_currency(a_qty), cell_right), Paragraph(format_currency(a_val), cell_right),
            Paragraph(format_currency(diff), cell_right)
        ])
        for c in range(7):
            qt_styles.append(('BACKGROUND', (c, idx), (c, idx), bg))
            
    tot_row = len(items) + 1
    qt_rows.append([
        Paragraph("<b>TỔNG CỘNG QUYẾT TOÁN</b>", cell_total), Paragraph("", cell_style),
        Paragraph("", cell_style), Paragraph(format_currency(tot_con_val), cell_right),
        Paragraph("", cell_style), Paragraph(format_currency(tot_act_val), cell_right),
        Paragraph(format_currency(tot_act_val - tot_con_val), cell_right)
    ])
    qt_styles.append(('SPAN', (0, tot_row), (2, tot_row)))
    qt_styles.append(('BACKGROUND', (0, tot_row), (-1, tot_row), colors.HexColor('#BDD7EE')))
    
    qt_table = Table(qt_rows, colWidths=[10*mm, 70*mm, 15*mm, 25*mm, 15*mm, 25*mm, 20*mm])
    qt_table.setStyle(TableStyle(qt_styles))
    story.append(qt_table)
    
    # ------------------ Section 5: BẢNG KÊ HÓA ĐƠN GTGT ------------------
    story.append(PageBreak())
    story.append(Paragraph("BẢNG KÊ HÓA ĐƠN GIÁ TRỊ GIA TĂNG (GTGT)", title_style))
    story.append(Spacer(1, 5))
    
    inv_headers = [
        Paragraph("STT", cell_header), Paragraph("Số hóa đơn", cell_header),
        Paragraph("Ngày hóa đơn", cell_header), Paragraph("Nhà cung cấp/Người bán", cell_header),
        Paragraph("Loại chi phí", cell_header), Paragraph("Trước thuế (VND)", cell_header),
        Paragraph("Thuế (VND)", cell_header), Paragraph("Tổng (VND)", cell_header)
    ]
    inv_rows = [inv_headers]
    inv_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D3D3D3')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]
    
    tot_pre_tax = 0
    tot_tax_val = 0
    tot_inv_val = 0
    for idx, inv in enumerate(invoices, 1):
        pre_tax = float(inv.get("truoc_thue") or 0)
        rate = float(inv.get("thue_suat") or 0.10)
        tax = pre_tax * rate
        total = pre_tax + tax
        tot_pre_tax += pre_tax
        tot_tax_val += tax
        tot_inv_val += total
        
        bg = colors.white if idx % 2 != 0 else colors.HexColor('#EBF3FB')
        inv_rows.append([
            Paragraph(str(idx), cell_center), Paragraph(inv.get("so_hd", ""), cell_center),
            Paragraph(inv.get("ngay", ""), cell_center), Paragraph(inv.get("ten_ncc", ""), cell_style),
            Paragraph(inv.get("loai", ""), cell_center), Paragraph(format_currency(pre_tax), cell_right),
            Paragraph(format_currency(tax), cell_right), Paragraph(format_currency(total), cell_right)
        ])
        for c in range(8):
            inv_styles.append(('BACKGROUND', (c, idx), (c, idx), bg))
            
    tot_row = len(invoices) + 1
    inv_rows.append([
        Paragraph("<b>TỔNG CỘNG HÓA ĐƠN</b>", cell_total), Paragraph("", cell_style),
        Paragraph("", cell_style), Paragraph("", cell_style), Paragraph("", cell_style),
        Paragraph(format_currency(tot_pre_tax), cell_right), Paragraph(format_currency(tot_tax_val), cell_right),
        Paragraph(format_currency(tot_inv_val), cell_right)
    ])
    inv_styles.append(('SPAN', (0, tot_row), (4, tot_row)))
    inv_styles.append(('BACKGROUND', (0, tot_row), (-1, tot_row), colors.HexColor('#BDD7EE')))
    
    inv_table = Table(inv_rows, colWidths=[10*mm, 20*mm, 20*mm, 55*mm, 20*mm, 20*mm, 15*mm, 20*mm])
    inv_table.setStyle(TableStyle(inv_styles))
    story.append(inv_table)
    
    # Build Document
    def add_page_decorations(canvas, doc):
        canvas.saveState()
        canvas.setFont(f"{FONT}-Italic" if FONT == 'Arial' else 'Helvetica-Oblique', 8)
        canvas.drawString(15*mm, 8*mm, f"Hồ sơ đề nghị thanh toán công trình | {billing_info.get('project_name', '')}")
        page_num = canvas.getPageNumber()
        canvas.drawRightString(210*mm - 15*mm, 8*mm, f"Trang {page_num}")
        canvas.restoreState()
        
    doc.build(story, onFirstPage=add_page_decorations, onLaterPages=add_page_decorations)

def build_invoices_csv(invoices, out_path):
    with open(out_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow([
            "stt", "so_hoa_don", "ngay_hoa_don", "ten_nha_cung_cap", "loai_chi_phi",
            "gia_tri_truoc_thue", "thue_suat", "tien_thue", "tong_tien_thanh_toan"
        ])
        for idx, inv in enumerate(invoices, 1):
            pre = float(inv.get("truoc_thue") or 0)
            rate = float(inv.get("thue_suat") or 0.10)
            tax = pre * rate
            writer.writerow([
                idx,
                inv.get("so_hd", ""),
                inv.get("ngay", ""),
                inv.get("ten_ncc", ""),
                inv.get("loai", "Vật tư"),
                pre,
                rate,
                tax,
                pre + tax
            ])

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python invoice_processor.py <payment_data_json> [project_id]")
        sys.exit(1)
        
    in_path = sys.argv[1]
    p_id = sys.argv[2] if len(sys.argv) > 2 else None
    
    res = generate_payment_package(in_path, p_id)
    print(json.dumps(res, ensure_ascii=False))
