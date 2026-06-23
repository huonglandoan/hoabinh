import os
import openpyxl
from datetime import datetime
from openpyxl.styles import PatternFill
from core_excel_utils import (
    apply_header_style,
    apply_data_style,
    apply_total_style,
    autofit_column_widths,
    FMT_CURRENCY,
    FMT_DATE,
    COLOR_ZEBRA_BG
)

def format_date_vietnam(date_str):
    if not date_str:
        return 'Chưa cập nhật'
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            dt = datetime.strptime(date_str.strip(), fmt)
            return dt.strftime("%d/%m/%Y")
        except ValueError:
            continue
    return date_str

def format_datetime_vietnam(dt_str):
    if not dt_str:
        return 'Chưa cập nhật'
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%dT%H:%M:%S"):
        try:
            dt = datetime.strptime(dt_str.strip(), fmt)
            return dt.strftime("%H:%M:%S - %d/%m/%Y")
        except ValueError:
            continue
    return dt_str

def build_workbook(quote_data, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bao_Gia_Chi_Tiet"
    
    # Hide grid lines? No, show grid lines
    ws.views.sheetView[0].showGridLines = True
    
    project_info = quote_data.get("project_info", {})
    items = quote_data.get("items", [])
    version = quote_data.get("version", "v1")
    generated_at = quote_data.get("generated_at", "")
    
    # Title Block
    ws['A1'] = "BẢNG BÁO GIÁ CHI TIẾT CÔNG TRÌNH"
    ws['A1'].font = openpyxl.styles.Font(name="Arial", size=16, bold=True, color="1F4E79")
    
    ws['A3'] = "Dự án:"
    ws['B3'] = project_info.get("project_name", "")
    ws['A4'] = "Địa điểm:"
    ws['B4'] = project_info.get("location", "")
    
    client_name = project_info.get("client_name") or project_info.get("clientName") or "Chưa cập nhật"
    contractor_name = project_info.get("contractor_name") or project_info.get("contractorName") or "Chưa cập nhật"
    contractor_name = contractor_name.replace("Cty ", "Công ty ")
    
    ws['A5'] = "Khách hàng (Bên A):"
    ws['B5'] = client_name
    ws['A6'] = "Đơn vị thi công (Bên B):"
    ws['B6'] = contractor_name
    
    start_dt = format_date_vietnam(project_info.get('start_date', ''))
    end_dt = format_date_vietnam(project_info.get('expected_completion_date', ''))
    gen_time = format_datetime_vietnam(generated_at)
    
    ws['A7'] = "Thời gian thi công:"
    ws['B7'] = f"{start_dt} đến {end_dt}"
    ws['A8'] = "Phiên bản:"
    ws['B8'] = f"{version} (Cập nhật lúc: {gen_time})"
    
    # Bold labels
    for row in range(3, 9):
        ws[f'A{row}'].font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
        ws[f'B{row}'].font = openpyxl.styles.Font(name="Arial", size=10)
        
    # Table Header at Row 10
    headers = [
        "STT", "Mã hạng mục", "Hạng mục công việc", "ĐVT", 
        "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)", "Trạng thái"
    ]
    
    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=10, column=col_idx)
        cell.value = text
        apply_header_style(cell)
        
    current_row = 11
    stt = 1
    
    # Color mapping for approval status
    fill_approved = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # green
    fill_pending = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # yellow
    fill_rejected = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # red
    fill_removed = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")  # gray
    
    for item in items:
        status = item.get("approval_status", "Chờ duyệt")
        bg_zebra = COLOR_ZEBRA_BG if stt % 2 == 0 else None
        
        # Write Item Details
        ws.cell(row=current_row, column=1, value=stt)
        ws.cell(row=current_row, column=2, value=item.get("item_code", ""))
        ws.cell(row=current_row, column=3, value=item.get("item_name", ""))
        ws.cell(row=current_row, column=4, value=item.get("unit", ""))
        ws.cell(row=current_row, column=5, value=item.get("quoted_quantity", 0))
        ws.cell(row=current_row, column=6, value=item.get("original_unit_price", 0))
        
        # Thành tiền formula: =E{row}*F{row}
        ws.cell(row=current_row, column=7, value=f"=E{current_row}*F{current_row}")
        ws.cell(row=current_row, column=8, value=status)
        
        # Apply data styles
        for col_idx in range(1, 9):
            cell = ws.cell(row=current_row, column=col_idx)
            apply_data_style(cell, bg_color=bg_zebra)
            if col_idx in (1, 2, 4, 8):
                cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
            elif col_idx in (5, 6, 7):
                cell.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
                if col_idx in (6, 7):
                    cell.number_format = FMT_CURRENCY
                        
        # Format status cell specifically
        status_cell = ws.cell(row=current_row, column=8)
        if status == "Đã phê duyệt":
            status_cell.fill = fill_approved
            status_cell.font = openpyxl.styles.Font(name="Arial", size=10, color="006100", bold=True)
        elif status == "Chờ duyệt":
            status_cell.fill = fill_pending
            status_cell.font = openpyxl.styles.Font(name="Arial", size=10, color="9C6500", bold=True)
        elif status == "Từ chối":
            status_cell.fill = fill_rejected
            status_cell.font = openpyxl.styles.Font(name="Arial", size=10, color="9C0006", bold=True)
        elif status == "Đã loại bỏ":
            status_cell.fill = fill_removed
            status_cell.font = openpyxl.styles.Font(name="Arial", size=10, color="555555", italic=True)
            
        current_row += 1
        stt += 1
        
    # Total Row
    total_row = current_row
    ws.merge_cells(f"A{total_row}:F{total_row}")
    total_label = ws.cell(row=total_row, column=1, value="TỔNG GIÁ TRỊ HỢP ĐỒNG (VND)")
    apply_total_style(total_label, align="right")
    
    total_val_cell = ws.cell(row=total_row, column=7, value=f"=SUM(G11:G{total_row-1})")
    apply_total_style(total_val_cell, align="right", num_format=FMT_CURRENCY)
    
    # Format rest of total row
    for col_idx in range(8, 9):
        c = ws.cell(row=total_row, column=col_idx)
        apply_total_style(c)
        
    # Set manual widths or autofit
    manual_widths = {
        'A': 5,   # STT
        'B': 15,  # Mã HM
        'C': 45,  # Hạng mục
        'D': 8,   # ĐVT
        'E': 12,  # SL
        'F': 16,  # Đơn giá
        'G': 18,  # Thành tiền
        'H': 15   # Trạng thái
    }
    autofit_column_widths(ws, padding=3, manual_widths=manual_widths)
    
    wb.save(out_path)

if __name__ == "__main__":
    import sys
    import json
    if len(sys.argv) < 3:
        print("Usage: python export_excel.py <quote_json_path> <out_xlsx_path>")
        sys.exit(1)
    with open(sys.argv[1], 'r', encoding='utf-8') as f:
        data = json.load(f)
    build_workbook(data, sys.argv[2])
    print(f"Excel exported to {sys.argv[2]}")
